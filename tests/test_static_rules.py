"""Issue #4: isolated predicates, candidate domains, and real PPTX integration."""

import json
from copy import deepcopy
from fractions import Fraction as F
from io import BytesIO
from itertools import permutations
from pathlib import Path
from typing import Any, cast
from xml.etree import ElementTree as ET

import pytest
from jsonschema import Draft202012Validator
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from pptviper import Bounds, Config, Report, inspect
from pptviper.cli import main
from pptviper.geometry import A
from pptviper.geometry import Interval as I
from pptviper.geometry_rules import connector, pair_rules, run_rules, unary_rules
from pptviper.inheritance import NS
from pptviper.package import P
from pptviper.report import measurement
from pptviper.rule_catalog import NEW_RULES, THRESHOLDS
from pptviper.rules import Item, Result, collect, emit, evaluate_rules, runs
from pptviper.text_rules import deck_rules, role_rules, text_rules
from tests.pptx_factory import replace_part


def item(
    key="a", box=(1000000, 1000000, 2000000, 2000000), text="", sid="256", kind="sp"
) -> Item:
    paragraphs = [
        {
            "properties": {},
            "children": {},
            "runs": [
                {
                    "text": text,
                    "kind": "r",
                    "properties": {
                        "sz": "1200",
                        "latin": "Arial",
                        "ea": "Noto Sans CJK",
                    },
                }
            ],
        }
    ]
    return Item(
        dict(id=sid, number=int(sid) - 255),
        dict(
            key=key,
            id=key,
            name=key,
            part=f"ppt/slides/slide{sid}.xml",
            kind=kind,
            origin="slide",
            hidden=False,
        ),
        dict(
            paragraphs=paragraphs,
            autofit={"mode": "noAutofit"},
            semantic=dict(
                text=text,
                preset="rect",
                placeholder=None,
                theme_fonts=["Arial"],
                layout_part="layout1",
                tables=[],
            ),
        ),
        cast(tuple[I, I, I, I], tuple(I.exact(v) for v in box)),
    )


def title(i: Item) -> Item:
    i.meta["placeholder"] = {"type": "title", "idx": "0"}
    return i


def paragraph(i: Item) -> dict[str, Any]:
    return i.props["paragraphs"][0]


def font(i: Item) -> dict[str, Any]:
    return paragraph(i)["runs"][0]["properties"]


def report_of(items: list[Item], config=None) -> Report:
    r = Report(config or Config())
    for i in items:
        row = dict(
            slide=i.slide,
            shapes=[i.ref],
            properties=i.props,
            locator=i.ref["key"],
            branch=None,
        )
        if i.bounds:
            row["bounds"] = {
                name: dict(low=measurement(v.low), high=measurement(v.high))
                for name, v in zip(
                    ("left", "top", "right", "bottom"), i.bounds, strict=True
                )
            }
        r.shapes.append(row)
    return r


def findings(report: Report, rule: str):
    return [f for f in report.findings if f["rule"] == rule]


def isolated(rule: str, hit: bool) -> Report:
    """Independent fixtures: one declared candidate for the target predicate."""
    a = item(text="Example")
    b = item("b", text="Other")
    result: Result
    members = [a]
    if rule in (
        "edge-margin",
        "zero-size",
        "tiny-shape",
        "hidden-shape",
        "layout-override",
    ):
        assert a.bounds is not None
        if rule == "edge-margin":
            a.bounds = (I.exact(457199 if hit else 457200), *a.bounds[1:])
        elif rule == "zero-size":
            a.bounds = (
                I.exact(1000000),
                I.exact(1000000),
                I.exact(1000000 if hit else 1000001),
                I.exact(2000000),
            )
        elif rule == "tiny-shape":
            a.bounds = cast(
                tuple[I, I, I, I],
                tuple(
                    I.exact(v)
                    for v in (1000000, 1000000, 1035999 if hit else 1036000, 1035999)
                ),
            )
        elif rule == "hidden-shape":
            a.ref["hidden"] = hit
        else:
            title(a).meta.update(
                local_geometry=True,
                layout_bounds={
                    k: [str(v), str(v)]
                    for k, v in zip(
                        ("left", "top", "right", "bottom"),
                        (1000000 + int(hit), 1000000, 2000000, 2000000),
                        strict=True,
                    )
                },
            )
        result = unary_rules(a, 10000000, 7500000, THRESHOLDS)[rule]
    elif rule in ("overlap", "too-close", "misaligned"):
        if rule == "overlap":
            b = item("b", (1999999 if hit else 2000000, 1000000, 3000000, 2000000))
        elif rule == "too-close":
            b = item("b", (2273599 if hit else 2273600, 1000000, 3000000, 2000000))
        else:
            b = item("b", (1000000 + int(hit), 1000000, 2000000 + int(hit), 2000000))
        members = [a, b]
        result = pair_rules(a, b, THRESHOLDS)[rule]
    elif rule in ("uneven-gap", "uneven-size"):
        a = item("a", (100, 100, 200, 200))
        b = item("b", (300, 100, 400 + int(hit and rule == "uneven-size"), 200))
        c = item("c", (500 + int(hit and rule == "uneven-gap"), 100, 600, 200))
        members = [a, b, c]
        result = next(res for name, _, res, _ in run_rules(members, 0) if name == rule)
    elif rule == "connector-unattached":
        a.ref["kind"] = "cxnSp"
        b.ref["part"] = a.ref["part"]
        a.meta.update(
            connections={}
            if hit
            else {end: dict(id="b", idx="0") for end in ("start", "end")},
            endpoints=[[["1500000", "1500000"], ["1000000", "1000000"]]] * 2,
        )
        result = connector(a, [a, b])
    elif rule in (
        "position-drift",
        "font-size-inconsistent",
        "line-spacing-inconsistent",
        "bullet-inconsistent",
    ):
        if rule == "position-drift":
            b = title(
                item(
                    "b",
                    (1000000 + int(hit), 1000000, 2000000, 2000000),
                    text="Other",
                    sid="257",
                )
            )
            title(a)
        elif rule == "font-size-inconsistent":
            font(b)["sz"] = "1300" if hit else "1200"
        elif rule == "line-spacing-inconsistent":
            paragraph(b)["children"]["lnSpc"] = dict(
                tag="lnSpc",
                attributes={},
                children=[
                    dict(
                        tag="spcPct",
                        attributes={"val": "120000" if hit else "100000"},
                        children=[],
                    )
                ],
            )
        elif hit:
            paragraph(b)["children"]["buChar"] = dict(
                tag="buChar", attributes={"char": "•"}, children=[]
            )
        members = [a, b]
        result = role_rules(a, b)[rule]
    elif rule in ("title-missing", "title-duplicate", "font-family-count"):
        r = Report()
        if rule == "title-missing":
            title(a).meta["text"] = "" if hit else "Title"
            deck_rules(r, [a])
        elif rule == "title-duplicate":
            a = title(a)
            b = title(item("b", text="Example" if hit else "Distinct", sid="257"))
            deck_rules(r, [a, b])
        else:
            font(b)["latin"] = "Verdana"
            c = item("c", text="Third")
            font(c)["latin"] = "Georgia" if hit else "Arial"
            deck_rules(r, [a, b, c])
        return r
    else:
        if rule == "autofit-shrunk":
            a.props["autofit"] = {
                "mode": "normAutofit",
                "fontScale": "99999" if hit else "100000",
            }
        elif rule == "font-too-small":
            font(a)["sz"] = "1199" if hit else "1200"
        elif rule == "east-asian-font-unset":
            a = item(text="\uff21")
            font(a)["ea"] = "" if hit else "Noto Sans CJK"
        elif rule == "empty-paragraph-spacing" and hit:
            a.props["paragraphs"].append(dict(runs=[]))
        elif rule == "manual-line-break" and hit:
            paragraph(a)["runs"].append(dict(kind="br", text=""))
        elif rule == "empty-placeholder":
            title(a).meta["inherited_text"] = ["Example" if hit else "Template"]
        elif rule == "text-hygiene":
            a.meta["text"] = "2026２０２６" if hit else "2026"
        elif rule == "table-column-uneven":
            a.meta["tables"] = [
                dict(widths=["100", "101" if hit else "100"], first_col=False)
            ]
        result = text_rules(a, THRESHOLDS)[rule]
    r = Report()
    emit(r, rule, members, result)
    return r


@pytest.mark.parametrize("rule", NEW_RULES)
@pytest.mark.parametrize("hit", [False, True])
def test_ac01_each_rule(rule, hit):
    r = isolated(rule, hit)
    assert len(findings(r, rule)) == int(hit)
    if hit:
        assert r.to_dict()["coverage"][rule]["finding"] == 1


@pytest.mark.parametrize(
    "threshold,delta,expected",
    [(72000, -72000, False), (72000, -1, True), (72000, 0, True), (72000, 1, False)],
)
def test_ac02_misaligned_boundaries(threshold, delta, expected):
    a = item(box=(0, 0, 1000000, 1000000))
    offset = threshold + delta
    b = item("b", box=(offset, offset, 1000000 + offset, 1000000 + offset))
    assert pair_rules(a, b, THRESHOLDS)["misaligned"].hit is expected


@pytest.mark.parametrize(
    "rule,limit",
    [("edge-margin", 457200), ("tiny-shape", 36000), ("too-close", 273600)],
)
@pytest.mark.parametrize("delta", [-1, 0, 1])
def test_ac02_strict_boundaries(rule, limit, delta):
    v = limit + delta
    a = item(box=(v, v, 1000000, 1000000))
    if rule == "tiny-shape":
        a = item(box=(1000000, 1000000, 1000000 + v, 1000000 + v))
    if rule == "too-close":
        a = item(box=(0, 0, 1000000, 1000000))
        b = item("b", box=(1000000 + v, 0, 2000000 + v, 1000000))
        result = pair_rules(a, b, THRESHOLDS)[rule]
    else:
        result = unary_rules(a, 10000000, 7500000, THRESHOLDS)[rule]
    assert result.hit is (delta < 0)


def test_ac02_interval_uncertainty():
    a = item()
    a.bounds = (
        I(F(457199), F(457201)),
        I.exact(1000000),
        I.exact(2000000),
        I.exact(2000000),
    )
    res = unary_rules(a, 10000000, 7500000, THRESHOLDS)["edge-margin"]
    assert not res.hit and res.unknown
    b = item("b")
    b.bounds = (
        I(F(1999999), F(2000001)),
        I.exact(1000000),
        I.exact(3000000),
        I.exact(2000000),
    )
    res = pair_rules(item(), b, THRESHOLDS)["overlap"]
    assert not res.hit and res.unknown


def test_ac03_touch_overlap_and_corner():
    a = item(box=(0, 0, 100, 100))
    for x, overlap, close in (
        (99, True, False),
        (100, False, True),
        (101, False, True),
    ):
        res = pair_rules(a, item("b", (x, 0, x + 100, 100)), THRESHOLDS)
        assert res["overlap"].hit is overlap
        assert res["too-close"].hit is close
    assert not pair_rules(a, item("b", (100, 100, 200, 200)), THRESHOLDS)[
        "too-close"
    ].hit


def test_ac04_runs_are_deterministic_nontransitive():
    items = [
        item("a", (0, 0, 100, 100)),
        item("b", (200, 50, 300, 150)),
        item("c", (401, 100, 501, 200)),
    ]
    for order in permutations(items):
        assert [[i.ref["id"] for i in run] for run in runs(list(order), 0)] == [
            ["a", "b"]
        ]
    items[2].bounds = (I.exact(401), I.exact(75), I.exact(501), I.exact(175))
    res = run_rules(runs(items, 0)[0], 0)[0][2]
    assert (
        res.hit
        and res.details == {"minimum_gap": "100", "maximum_gap": "101"}
        and res.value == 1
    )


def test_ac05_cross_slide_identity_and_role_union():
    a = title(item(text="Title"))
    b = title(item("b", (1000001, 1000000, 2000000, 2000000), text="Other", sid="257"))

    def run(items):
        r = report_of(items)
        evaluate_rules(r, 10000000, 7500000)
        return r

    r = run([a, b])
    first = findings(r, "position-drift")[0]
    assert first["measured"]["value"] == 1
    assert {s["slide_id"] for s in first["shapes"]} == {"256", "257"}
    assert b.bounds is not None
    b.bounds = (I.exact(1000002), *b.bounds[1:])
    assert (
        first["fingerprint"]
        == findings(run([b, a]), "position-drift")[0]["fingerprint"]
    )
    assert r.to_dict()["coverage"]["font-size-inconsistent"]["candidates"] == 1
    b.bounds = a.bounds
    assert not findings(run([a, b]), "position-drift")


def test_ac06_mixed_width_reasons_aggregated():
    r = isolated("text-hygiene", True)
    assert findings(r, "text-hygiene")[0]["details"]["reasons"] == [
        "mixed full/half-width Latin letters or digits"
    ]
    a = item(text=" 2026２０２６ ((")
    res = text_rules(a, THRESHOLDS)["text-hygiene"]
    assert len(res.details["reasons"]) >= 3


def test_ac07_accounting_all_domains_and_unknown_finding():
    items = [
        title(item("a", (100, 100, 200, 200), "A")),
        title(item("b", (300, 100, 400, 200), "B")),
        title(item("c", (501, 100, 601, 200), "C")),
        title(item("d", (100, 100, 200, 200), "A", sid="257")),
    ]
    r = report_of(items)
    evaluate_rules(r, 10000000, 7500000)
    coverage = r.to_dict()["coverage"]
    for rule in NEW_RULES:
        c = coverage[rule]
        assert c["candidates"] == sum(
            c[k] for k in ("pass", "finding", "unknown", "not-applicable")
        )
    assert coverage["overlap"]["candidates"] == 3
    assert coverage["uneven-gap"]["candidates"] == 1
    assert coverage["font-size-inconsistent"]["candidates"] == 6
    assert coverage["title-missing"]["candidates"] == 2
    assert coverage["title-duplicate"]["candidates"] == 3
    assert coverage["font-family-count"]["candidates"] == 1
    a = item(text="A")
    font(a).pop("sz")
    paragraph(a)["runs"].append(dict(text="B", properties={"sz": "1100"}))
    result = text_rules(a, THRESHOLDS)["font-too-small"]
    assert result.hit and result.unknown
    r = Report()
    emit(r, "font-too-small", [a], result)
    assert (
        list(r.candidates["font-too-small"].values()) == ["finding"]
        and len(r.unknown) == 1
    )


def test_ac08_hierarchy_hidden_and_alternatives():
    items = [
        item("a"),
        item("b"),
        item("group", kind="grpSp"),
        item("cell", kind="table-cell"),
    ]
    items[1].ref["hidden"] = True
    r = report_of(items)
    evaluate_rules(r, 10000000, 7500000)
    assert r.to_dict()["coverage"]["overlap"]["candidates"] == 1
    assert len(findings(r, "hidden-shape")) == 1
    r = report_of([items[0], deepcopy(items[0]), items[1]])
    assert len(collect(r)) == 2 and collect(r)[0].ambiguous
    evaluate_rules(r, 10000000, 7500000)
    assert not findings(r, "overlap")
    assert r.to_dict()["coverage"]["overlap"]["unknown"] == 1


def deck_bytes(*, empty=False) -> bytes:
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    if not empty:
        sh = s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Emu(457199), Emu(1000000), Emu(1000000), Emu(1000000)
        )
        sh.text = "Example"
        sh.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
        sh.text_frame.paragraphs[0].runs[0].font.name = "Arial"
    out = BytesIO()
    p.save(out)
    return out.getvalue()


def test_ac09_cli_config_schema_and_determinism(tmp_path, capsys):
    data = deck_bytes()
    pptx = tmp_path / "input.pptx"
    pptx.write_bytes(data)
    settings = tmp_path / "settings.toml"
    settings.write_text(
        "[thresholds]\nedge-margin = 457199\nfont-too-small = 10\n"
        '[[suppress]]\nrule = "title-missing"\n'
    )
    r = inspect(data, config=settings)
    assert not findings(r, "edge-margin") and not findings(r, "font-too-small")
    assert findings(r, "title-missing")[0]["suppressed"]
    assert r.to_json() == inspect(data, config=settings).to_json()
    assert main([str(pptx), "--config", str(settings)]) == r.exit_code
    assert json.loads(capsys.readouterr().out) == r.to_dict()
    schema = json.loads(
        (Path(__file__).parents[1] / "src/pptviper/report.schema.json").read_text()
    )
    Draft202012Validator(schema).validate(r.to_dict())
    Draft202012Validator(schema).validate(inspect(data).to_dict())
    for rule in NEW_RULES:
        # All positive rule outputs conform to the extended metric/ref schema.
        Draft202012Validator(schema).validate(isolated(rule, True).to_dict())


@pytest.mark.parametrize(
    "settings",
    [
        "[thresholds]\nunknown=1",
        "[thresholds]\nedge-margin=-1",
        "[thresholds]\nedge-margin=nan",
        "[thresholds]\nedge-margin=inf",
        "[thresholds]\nedge-margin=true",
        "[thresholds]\nfont-family-count=2.5",
        '[[suppress]]\nrule="unknown"',
    ],
)
def test_ac09_invalid_settings(tmp_path, settings):
    config = tmp_path / "bad.toml"
    config.write_text(settings)
    r = inspect(deck_bytes(), config=config)
    assert r.exit_code == 3 and r.errors


def test_ac09_cross_slide_suppression():
    a = title(item(text="Title"))
    b = title(item("b", text="Title", sid="257"))
    r = Report(Config(suppress=[dict(rule="title-duplicate", slide="257")]))
    deck_rules(r, [a, b])
    assert findings(r, "title-duplicate")[0]["suppressed"]


def test_ac10_missing_properties_and_connector_proxy():
    a = item(text="Example")
    a.bounds = None
    font(a).pop("sz")
    assert unary_rules(a, 10000000, 7500000, THRESHOLDS)["edge-margin"].unknown
    assert text_rules(a, THRESHOLDS)["font-too-small"].unknown
    a.meta["text_unknown"] = ["Missing inheritance"]
    a.meta["text"] = "\uff21"
    paragraph(a)["runs"][0]["text"] = "\uff21"
    font(a)["ea"] = ""
    res = text_rules(a, THRESHOLDS)["east-asian-font-unset"]
    assert res.hit and res.unknown
    c = item("c", kind="cxnSp")
    target = item("t")
    target.ref["part"] = c.ref["part"]
    c.meta.update(
        connections={end: dict(id="t", idx="7") for end in ("start", "end")},
        endpoints=[[["1500001", "1500001"], ["1000000", "1000000"]]] * 2,
    )
    res = connector(c, [c, target])
    assert res.hit and res.unknown and res.value == 1
    r = Report()
    emit(r, "connector-unattached", [c], res)
    assert findings(r, "connector-unattached")[0]["confidence"] == "worst-case"
    c.meta["endpoints"] = [[["1500000", "1500000"], ["1000000", "1000000"]]] * 2
    res = connector(c, [c, target])
    assert not res.hit and res.unknown
    c.bounds = None
    assert connector(c, [c, target]).unknown


def test_ac11_expanded_bounds_and_empty_slide():
    data = deck_bytes()

    def provider(context, bounds):
        return Bounds(I.exact(-1), bounds.top, bounds.right, bounds.bottom), []

    r = inspect(data, bounds_provider=provider)
    assert findings(r, "edge-margin")[0]["measured"]["value"] == -1
    assert len(findings(inspect(deck_bytes(empty=True)), "title-missing")) == 1


def test_saved_ooxml_metadata_and_text_rules():
    data = deck_bytes()
    from zipfile import ZipFile

    with ZipFile(BytesIO(data)) as z:
        root = ET.fromstring(z.read("ppt/slides/slide1.xml"))
    shape = root.find("p:cSld/p:spTree/p:sp", NS)
    assert shape is not None
    nv = shape.find("p:nvSpPr/p:cNvPr", NS)
    assert nv is not None
    nv.set("hidden", "1")
    nvpr = shape.find("p:nvSpPr/p:nvPr", NS)
    assert nvpr is not None
    ET.SubElement(nvpr, P + "ph", {"type": "title", "idx": "0"})
    body = shape.find("p:txBody/a:bodyPr", NS)
    assert body is not None
    for child in list(body):
        body.remove(child)
    ET.SubElement(body, A + "normAutofit", {"fontScale": "80000"})
    para = shape.find("p:txBody/a:p", NS)
    assert para is not None
    ET.SubElement(para, A + "br")
    tx = shape.find("p:txBody", NS)
    assert tx is not None
    ET.SubElement(tx, A + "p")
    r = inspect(replace_part(data, "ppt/slides/slide1.xml", ET.tostring(root)))
    for rule in (
        "hidden-shape",
        "autofit-shrunk",
        "manual-line-break",
        "empty-paragraph-spacing",
        "font-too-small",
    ):
        assert len(findings(r, rule)) == 1
    assert not findings(r, "title-missing")
    assert findings(r, "font-too-small")[0]["measured"]["exact_value"] == "44/5"


def test_spacing_bullet_defaults_and_font_sets():
    a = item(text="A")
    b = item("b", text="B")
    paragraph(b)["children"]["buNone"] = dict(tag="buNone", attributes={}, children=[])
    assert not role_rules(a, b)["bullet-inconsistent"].hit
    paragraph(a)["runs"].append(dict(text="C", properties={"sz": "1300"}))
    paragraph(b)["runs"].insert(0, dict(text="D", properties={"sz": "1300"}))
    assert not role_rules(a, b)["font-size-inconsistent"].hit
    font(b)["sz"] = "1400"
    assert role_rules(a, b)["font-size-inconsistent"].hit


def test_note_font_threshold_and_table_first_column():
    a = item(text="Footnote", kind="table-cell")
    for sz, expected in (("899", True), ("900", False), ("901", False)):
        font(a)["sz"] = sz
        assert text_rules(a, THRESHOLDS)["font-too-small"].hit is expected
    a.meta["tables"] = [dict(widths=["100", "200", "200"], first_col=True)]
    assert not text_rules(a, THRESHOLDS)["table-column-uneven"].hit
    a.meta["tables"][0]["widths"] = ["100", "200"]
    assert not text_rules(a, THRESHOLDS)["table-column-uneven"].applicable


def test_real_slide_roles_tables_and_connectors():
    from pptx.enum.shapes import MSO_CONNECTOR

    p = Presentation()
    first = p.slides.add_slide(p.slide_layouts[0])
    second = p.slides.add_slide(p.slide_layouts[0])
    left_title, right_title = first.shapes.title, second.shapes.title
    assert left_title is not None and right_title is not None
    left_title.text = right_title.text = "Title"
    left_title.left = Emu(right_title.left + 1)
    left_title.top = right_title.top
    left_title.width = right_title.width
    left_title.height = right_title.height
    table = first.shapes.add_table(
        1, 3, Emu(1000000), Emu(3000000), Emu(3000000), Emu(500000)
    ).table
    table.first_col = True
    table.columns[1].width = Emu(1000000)
    table.columns[2].width = Emu(1000001)
    first.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Emu(1000000), Emu(5000000), Emu(2000000), Emu(5000000)
    )
    out = BytesIO()
    p.save(out)
    r = inspect(out.getvalue())
    for rule in (
        "title-duplicate",
        "position-drift",
        "table-column-uneven",
        "connector-unattached",
        "layout-override",
    ):
        assert len(findings(r, rule)) == 1
    assert not findings(r, "zero-size")
    assert findings(r, "connector-unattached")[0]["confidence"] == "certain"
    assert not findings(r, "title-missing")


def test_incomplete_input_aggregate_and_alternative_presence():
    r = Report()
    r.slide_contexts = {"256": dict(slide=dict(id="256", number=1), complete=False)}
    r.unenumerated = [dict(part="missing.xml", reason="Unreadable input")]
    deck_rules(r, [])
    assert r.to_dict()["coverage"]["title-missing"]["unknown"] == 1
    assert r.to_dict()["coverage"]["font-family-count"]["unknown"] == 1
    assert not r.findings
    r = report_of([title(item(text="Title"))])
    r.shapes[0]["branch"] = "choice/0"
    evaluate_rules(r, 10000000, 7500000)
    assert r.to_dict()["coverage"]["title-missing"]["unknown"] == 1
    assert not findings(r, "title-missing")
    assert any(u["rule"] == "title-duplicate" for u in r.unknown)


def test_invalid_spacing_and_autofit_are_unknown():
    a = item(text="A")
    b = item("b", text="B")
    paragraph(a)["children"]["lnSpc"] = dict(
        children=[dict(tag="spcPct", attributes={"val": "bad"})]
    )
    assert role_rules(a, b)["line-spacing-inconsistent"].unknown
    a.props["autofit"] = {"mode": "normAutofit", "fontScale": "invalid"}
    assert text_rules(a, THRESHOLDS)["autofit-shrunk"].unknown
    assert text_rules(a, THRESHOLDS)["font-too-small"].unknown


def test_inherited_bullet_choice_override():
    from pptviper.inheritance import resolve

    shape = ET.fromstring(
        f'<p:sp xmlns:p="{P[1:-1]}" xmlns:a="{A[1:-1]}">'
        '<p:txBody><a:lstStyle><a:lvl1pPr><a:buChar char="•"/>'
        "</a:lvl1pPr></a:lstStyle><a:p><a:pPr><a:buNone/></a:pPr>"
        "<a:r><a:t>Text</a:t></a:r></a:p></p:txBody></p:sp>"
    )
    props, _ = resolve([shape], None, ET.Element(P + "presentation"), None, {})
    children = props["paragraphs"][0]["children"]
    assert "buNone" in children and "buChar" not in children


def test_normalized_zero_spacing_and_bullet_defaults():
    a = item(text="A")
    b = item("b", text="B")
    paragraph(b)["children"]["spcBef"] = dict(
        children=[dict(tag="spcPct", attributes={"val": "0"})]
    )
    assert not role_rules(a, b)["line-spacing-inconsistent"].hit
    for i in (a, b):
        paragraph(i)["children"]["buChar"] = dict(
            tag="buChar", attributes={"char": "•"}, children=[]
        )
    paragraph(b)["children"]["buFontTx"] = dict(
        tag="buFontTx", attributes={}, children=[]
    )
    paragraph(b)["children"]["buSzPct"] = dict(
        tag="buSzPct", attributes={"val": "100000"}, children=[]
    )
    assert not role_rules(a, b)["bullet-inconsistent"].hit
