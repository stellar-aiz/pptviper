"""Issue #3 acceptance fixtures and independent expected results."""

import json
import subprocess
from copy import deepcopy
from fractions import Fraction as F
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from jsonschema import Draft202012Validator
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu

from pptviper import Bounds, Config
from pptviper import inspect as inspect_all
from pptviper.cli import main
from pptviper.geometry import A, Interval, Matrix, transform, trig
from pptviper.inheritance import NS, color, resolve_table
from pptviper.package import P
from pptviper.report import Report
from pptviper.rule_catalog import NEW_RULES
from tests.pptx_factory import replace_part, single_shape_pptx

SLIDE = "ppt/slides/slide1.xml"
MASTER = "ppt/slideMasters/slideMaster1.xml"
LAYOUT = "ppt/slideLayouts/slideLayout7.xml"


def fixture(x: int = 0, y: int = 100, width: int = 200, height: int = 300) -> bytes:
    return single_shape_pptx(x=x, y=y, width=width, height=height)


def root(data: bytes, part: str = SLIDE) -> ET.Element:
    with ZipFile(BytesIO(data)) as z:
        return ET.fromstring(z.read(part))


def save(data: bytes, node: ET.Element, part: str = SLIDE) -> bytes:
    return replace_part(data, part, ET.tostring(node))


def shape(node: ET.Element) -> ET.Element:
    value = required(node, "p:cSld/p:spTree/p:sp", NS)
    assert value is not None
    return value


def tx(x: int = 0, y: int = 0, w: int = 100, h: int = 100, rot: int = 0) -> ET.Element:
    return ET.fromstring(
        f'''<a:xfrm xmlns:a="{A[1:-1]}" rot="{rot}">
<a:off x="{x}" y="{y}"/>
<a:ext cx="{w}" cy="{h}"/>
</a:xfrm>'''
    )


def findings(report: Report, rule: str = "off-slide") -> list:
    return [f for f in report.findings if f["rule"] == rule]


def test_ac01_exact_boundary() -> None:
    clean = inspect(fixture())
    assert clean.exit_code == 0
    result = inspect(fixture(-1))
    assert len(findings(result)) == 1
    assert findings(result)[0]["measured"]["emu"] == 1
    assert findings(result)[0]["confidence"] == "certain"


def test_ac02_placeholder_asymmetric_inheritance() -> None:
    data = fixture()
    slide, layout, master = root(data), root(data, LAYOUT), root(data, MASTER)
    s = shape(slide)
    nv = required(s, "p:nvSpPr/p:nvPr", NS)
    assert nv is not None
    ET.SubElement(nv, P + "ph", idx="42", type="obj")
    sp = required(s, P + "spPr")
    assert sp is not None
    sp.clear()
    layout_shape = deepcopy(s)
    required(layout_shape, "p:nvSpPr/p:cNvPr", NS).set("id", "88")
    required(layout_shape, "p:nvSpPr/p:nvPr/p:ph", NS).set("type", "subTitle")
    required(layout, "p:cSld/p:spTree", NS).append(layout_shape)
    m = next(
        e
        for e in master.findall("p:cSld/p:spTree/p:sp", NS)
        if required(e, "p:nvSpPr/p:nvPr/p:ph", NS).get("type") == "body"
    )
    required(m, "p:nvSpPr/p:nvPr/p:ph", NS).set("idx", "999")
    required(m, P + "spPr").clear()
    required(m, P + "spPr").append(tx(-1, 0))
    for part, node in ((SLIDE, slide), (LAYOUT, layout), (MASTER, master)):
        data = save(data, node, part)
    r = inspect(data)
    assert len(findings(r)) == 1
    assert findings(r)[0]["measured"]["emu"] == 1
    assert r.to_dict()["coverage"]["off-slide"]["candidates"] == 1


def test_ac03_cardinal_and_arbitrary_geometry() -> None:
    b, _ = transform(tx(10, 20, 100, 200, 5400000), Matrix())
    assert (b.left.low, b.top.low, b.right.high, b.bottom.high) == (-40, 70, 160, 170)
    t = tx(10, 20, 100, 200)
    t.set("flipH", "1")
    b, _ = transform(t, Matrix())
    assert (b.left.low, b.right.high) == (10, 110)
    s, c = trig(2700000)
    # Independently calculated sqrt(2)/2 at 60 decimal digits.
    reference = F("0.707106781186547524400844362104849039284835937688474036588340")
    assert s.low < reference < s.high and c.low < reference < c.high
    matrix = Matrix()
    for _ in range(3):
        g = tx(1, 2, 50, 150)
        ET.SubElement(g, A + "chOff", x="0", y="0")
        ET.SubElement(g, A + "chExt", cx="100", cy="100")
        _, matrix = transform(g, matrix, True)
    b, _ = transform(tx(8, 8, 8, 8), matrix)
    assert b.left.low == F(11, 4)
    assert b.top.low == F(73, 2)
    assert b.right.high == F(15, 4)


@pytest.mark.parametrize("offset, expected", [(-8, 1), (-1, 0.125)])
def test_ac03_nested_pptx(offset: int, expected: float) -> None:
    data = fixture(offset, 0, 8, 8)
    slide = root(data)
    tree = required(slide, "p:cSld/p:spTree", NS)
    s = shape(slide)
    tree.remove(s)
    for i in range(3):
        g = ET.Element(P + "grpSp")
        nv = ET.SubElement(g, P + "nvGrpSpPr")
        ET.SubElement(nv, P + "cNvPr", id=str(10 + i), name="group")
        gp = ET.SubElement(g, P + "grpSpPr")
        t = tx(0, 0, 50, 50)
        ET.SubElement(t, A + "chOff", x="0", y="0")
        ET.SubElement(t, A + "chExt", cx="100", cy="100")
        gp.append(t)
        g.append(s)
        s = g
    tree.append(s)
    r = inspect(save(data, slide))
    assert len(findings(r)) == 1
    assert findings(r)[0]["measured"]["emu"] == expected
    assert r.to_dict()["coverage"]["off-slide"]["candidates"] == 4


def test_ac04_master_logo_and_visibility() -> None:
    data = fixture()
    master = root(data, MASTER)
    logo = deepcopy(shape(root(data)))
    required(logo, "p:spPr/a:xfrm/a:off", NS).set("x", "-1")
    required(logo, "p:nvSpPr/p:cNvPr", NS).set("id", "99")
    required(master, "p:cSld/p:spTree", NS).append(logo)
    data = save(data, master, MASTER)
    r = inspect(data)
    assert findings(r)[0]["shapes"][0]["origin"] == "master"
    slide = root(data)
    slide.set("showMasterSp", "0")
    assert not findings(inspect(save(data, slide)))


def test_ac05_hidden() -> None:
    data = fixture(-1)
    slide = root(data)
    slide.set("show", "0")
    required(shape(slide), "p:nvSpPr/p:cNvPr", NS).set("hidden", "1")
    r = inspect(save(data, slide))
    f = findings(r)[0]
    assert f["slide"]["hidden"] and f["shapes"][0]["hidden"]
    assert r.to_dict()["coverage"]["off-slide"]["finding"] == 1


@pytest.mark.parametrize("uri", ["chart", "diagram", "unrecognized"])
def test_ac06_unsupported(uri: str) -> None:
    data = fixture()
    slide = root(data)
    s = shape(slide)
    s.tag = P + "graphicFrame"
    t = required(s, "p:spPr/a:xfrm", NS)
    required(s, P + "spPr").remove(t)
    t.tag = P + "xfrm"
    s.append(t)
    graphic = ET.SubElement(s, A + "graphic")
    ET.SubElement(graphic, A + "graphicData", uri=uri)
    r = inspect(save(data, slide))
    assert r.exit_code == 2 and r.unknown
    assert r.to_dict()["coverage"]["off-slide"]["unknown"] == 1


def test_ac07_alternatives() -> None:
    data = fixture(-1)
    slide = root(data)
    tree = required(slide, "p:cSld/p:spTree", NS)
    s = shape(slide)
    tree.remove(s)
    mc = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
    ac = ET.SubElement(tree, mc + "AlternateContent")
    ET.SubElement(ac, mc + "Choice", Requires="p14").append(s)
    other = deepcopy(s)
    required(other, "p:spPr/a:xfrm/a:off", NS).set("x", "-9")
    ET.SubElement(ac, mc + "Fallback").append(other)
    r = inspect(save(data, slide))
    assert len(findings(r)) == 1 and findings(r)[0]["measured"]["emu"] == 9
    assert r.to_dict()["coverage"]["off-slide"]["candidates"] == 1
    ET.SubElement(other, P + "extLst")
    assert inspect(save(data, slide)).unknown


@pytest.mark.parametrize("mode", ["zip", "xml", "missing", "size", "count"])
def test_ac08_unreadable(mode: str) -> None:
    data = fixture()
    config = Config()
    if mode == "zip":
        data = data[:20]
    elif mode == "xml":
        data = replace_part(data, SLIDE, b"<bad")
    elif mode == "missing":
        data = replace_part(data, SLIDE, None)
    elif mode == "size":
        config = Config(max_expanded_bytes=100)
    else:
        config = Config(max_parts=1)
    r = inspect(data, config=config)
    assert r.exit_code == 3 and findings(r, "unreadable-part")
    assert r.unenumerated


def test_ac08_partial_slides() -> None:
    p = Presentation()
    for _ in range(2):
        s = p.slides.add_slide(p.slide_layouts[6])
        s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(-1), Emu(0), Emu(10), Emu(10))
    output = BytesIO()
    p.save(output)
    data = replace_part(output.getvalue(), SLIDE, b"<bad")
    r = inspect(data)
    assert r.exit_code == 3
    assert findings(r)[0]["slide"]["number"] == 2


def test_ac09_determinism_and_movement() -> None:
    data = fixture(-1)
    assert inspect(data).to_json() == inspect(data).to_json()
    assert (
        findings(inspect(data))[0]["fingerprint"]
        == findings(inspect(fixture(-2)))[0]["fingerprint"]
    )


def test_ac09_slide_order() -> None:
    p = Presentation()
    for _ in range(2):
        s = p.slides.add_slide(p.slide_layouts[6])
        s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(-1), Emu(0), Emu(10), Emu(10))
    stream = BytesIO()
    p.save(stream)
    data = stream.getvalue()
    before = {f["fingerprint"] for f in findings(inspect(data))}
    node = root(data, "ppt/presentation.xml")
    ids = required(node, P + "sldIdLst")
    ids[:] = list(reversed(ids))
    after = inspect(save(data, node, "ppt/presentation.xml"))
    assert before == {f["fingerprint"] for f in findings(after)}


def test_ac10_schema(tmp_path: Path) -> None:
    schema = json.loads(
        (Path(__file__).parents[1] / "src/pptviper/report.schema.json").read_text()
    )
    Draft202012Validator.check_schema(schema)
    validator = Draft202012Validator(schema)
    config = tmp_path / "invalid.toml"
    config.write_text("unknown = true")
    for report in [
        inspect(fixture()),
        inspect(fixture(-1)),
        inspect(b"broken"),
        inspect(fixture(), config=config),
        inspect(replace_part(fixture(), SLIDE, b"<bad")),
    ]:
        validator.validate(json.loads(report.to_json()))
    bad = inspect(fixture()).to_dict()
    bad["exit_code"] = 99
    assert list(validator.iter_errors(bad))


@pytest.mark.parametrize(
    "selector", [{"rule": "off-slide"}, {"slide": "256"}, {"shape": SLIDE + "#2"}]
)
def test_ac11_suppression(selector: dict[str, str]) -> None:
    r = inspect(fixture(-1), config=Config(suppress=[selector]))
    assert r.exit_code == 0 and findings(r)[0]["suppressed"]
    assert r.to_dict()["coverage"]["off-slide"]["finding"] == 1
    assert r.to_dict()["effective_config"]["suppress"][0] == selector


@pytest.mark.parametrize(
    "content",
    [
        "unknown = 1",
        "max_parts = true",
        "max_parts = 0",
        "off_slide_threshold_emu = -1",
        '[[suppress]]\nrule="not-a-rule"',
        "[[suppress]]\nslide=256",
        'suppress="bad"',
    ],
)
def test_ac12_bad_config(content: str, tmp_path: Path) -> None:
    path = tmp_path / "config.toml"
    path.write_text(content)
    r = inspect(fixture(), config=path)
    assert r.exit_code == 3 and r.errors[0]["kind"] == "configuration"
    assert not findings(r, "unreadable-part")


def test_ac13_priority() -> None:
    r = Report(Config(suppress=[{"rule": "off-slide"}]))
    r.finding("off-slide", "s", "outside")
    assert r.exit_code == 0
    r.unresolved("unknown")
    assert r.exit_code == 2
    r.finding("invalid-ooxml", "part", "malformed")
    assert r.exit_code == 1
    r.finding("unreadable-part", "broken", "cannot read")
    assert r.exit_code == 3
    assert (
        inspect(
            b"broken", config=Config(suppress=[{"rule": "unreadable-part"}])
        ).exit_code
        == 3
    )


def test_ac14_no_io_side_effects(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    data = fixture()
    path = tmp_path / "input.pptx"
    path.write_bytes(data)

    def forbidden(*args, **kwargs):
        raise AssertionError("External process attempted")

    monkeypatch.setattr(subprocess, "Popen", forbidden)
    assert inspect(path).exit_code == 0
    assert path.read_bytes() == data
    # pytest-socket blocks all sockets for every test in this module.


def test_ac15_fractional_and_uncertain() -> None:
    def fractional(info, bounds):
        return Bounds(
            Interval.exact(F(-1, 2)), bounds.top, bounds.right, bounds.bottom
        ), []

    r = inspect(fixture(), bounds_provider=fractional)
    assert findings(r)[0]["measured"]["emu"] == 0.5

    def uncertain(info, bounds):
        return Bounds(
            Interval(F(-1, 10**30), F(1, 10**30)),
            bounds.top,
            bounds.right,
            bounds.bottom,
        ), []

    r = inspect(fixture(), bounds_provider=uncertain)
    assert r.exit_code == 2 and not findings(r)
    assert "Numerical enclosure" in r.unknown[0]["reason"]


def test_ac16_invalid_geometry_and_entities() -> None:
    data = fixture()
    node = root(data)
    required(shape(node), "p:spPr/a:xfrm/a:ext", NS).set("cx", "bad")
    r = inspect(save(data, node))
    assert findings(r, "invalid-ooxml") and r.unknown and not findings(r)
    malicious = (
        b'<!DOCTYPE x [<!ENTITY file SYSTEM "file:///etc/passwd">]><x>&file;</x>'
    )
    r = inspect(replace_part(data, SLIDE, malicious))
    assert r.exit_code == 3 and "root:" not in r.to_json()


def test_ac17_accounting() -> None:
    def provider(info, bounds):
        return Bounds.rectangle(-1, 0, 10, 10), ["Text extent remains uncertain"]

    r = inspect(
        fixture(),
        bounds_provider=provider,
        config=Config(suppress=[{"rule": "off-slide"}]),
    )
    counts = r.to_dict()["coverage"]["off-slide"]
    assert counts["candidates"] == counts["finding"] == 1
    assert counts["unknown"] == 0 and r.exit_code == 2
    r.candidates["off-slide"]["corrupt"] = "invalid"
    with pytest.raises(AssertionError, match="accounting"):
        r.to_dict()


def test_ac18_identity_and_multi_edge() -> None:
    r = inspect(fixture(-1, -2))
    assert len(findings(r)) == 1 and findings(r)[0]["measured"]["emu"] == 2
    data = fixture(-1)
    node = root(data)
    s = shape(node)
    tree = required(node, "p:cSld/p:spTree", NS)
    tree.append(deepcopy(s))
    r = inspect(save(data, node))
    assert len({f["fingerprint"] for f in findings(r)}) == 2
    assert findings(r, "invalid-ooxml")
    for n in tree.findall(P + "sp"):
        del required(n, "p:nvSpPr/p:cNvPr", NS).attrib["id"]
    assert inspect(save(data, node)).to_json() == inspect(save(data, node)).to_json()


def test_ac19_stored_properties_and_colors() -> None:
    data = fixture()
    node = root(data)
    s = shape(node)
    body = required(s, "p:txBody/a:bodyPr", NS)
    body.set("lIns", "123")
    ET.SubElement(body, A + "normAutofit", fontScale="80000", lnSpcReduction="20000")
    fill = ET.SubElement(required(s, P + "spPr"), A + "solidFill")
    clr = ET.SubElement(fill, A + "schemeClr", val="accent1")
    ET.SubElement(clr, A + "alpha", val="50000")
    props = inspect(save(data, node)).shapes[0]["properties"]
    assert props["body"]["lIns"] == "123"
    assert props["autofit"]["fontScale"] == "80000"
    assert props["fill"]["colors"][0]["alpha"] == 0.5
    unknown = []
    c = ET.fromstring(
        f'''<a:srgbClr xmlns:a="{A[1:-1]}" val="808080">
<a:lumMod val="50000"/>
<a:lumOff val="25000"/>
<a:tint val="50000"/>
<a:shade val="50000"/>
</a:srgbClr>'''
    )
    value = color(c, {}, {}, unknown)
    assert value and value["rgb"] == pytest.approx([95.6875] * 3) and not unknown


def test_ac19_table_styles() -> None:
    xml = f'''<a:tbl xmlns:a="{A[1:-1]}">
<a:tblPr firstRow="1">
<a:tableStyleId>custom</a:tableStyleId>
</a:tblPr>
<a:tr>
<a:tc>
<a:tcPr/>
</a:tc>
</a:tr>
</a:tbl>'''
    styles = ET.fromstring(
        f'''<a:tblStyleLst xmlns:a="{A[1:-1]}">
<a:tblStyle styleId="custom">
<a:wholeTbl>
<a:tcStyle>
<a:fill>
<a:solidFill>
<a:srgbClr val="000000"/>
</a:solidFill>
</a:fill>
</a:tcStyle>
</a:wholeTbl>
<a:firstRow>
<a:tcStyle>
<a:fill>
<a:solidFill>
<a:srgbClr val="ffffff"/>
</a:solidFill>
</a:fill>
</a:tcStyle>
</a:firstRow>
</a:tblStyle>
</a:tblStyleLst>'''
    )
    unknown = []
    r = resolve_table(ET.fromstring(xml), styles, {}, {}, unknown)
    assert r["cells"][0]["fill"]["rgb"] == [255, 255, 255]
    assert r["cells"][0]["sources"] == ["wholeTbl", "firstRow"] and not unknown


def test_cli_outputs(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    path = tmp_path / "input.pptx"
    path.write_bytes(fixture(-1))
    assert main([str(path)]) == 1
    assert json.loads(capsys.readouterr().out)["exit_code"] == 1
    assert main([str(path), "--format", "text"]) == 1
    assert "off-slide" in capsys.readouterr().out


def required(
    node: ET.Element, path: str, ns: dict[str, str] | None = None
) -> ET.Element:
    result = node.find(path, ns)
    assert result is not None, path
    return result


def test_duplicate_zip_is_structured() -> None:
    data = fixture()
    stream = BytesIO(data)
    with pytest.warns(UserWarning, match="Duplicate"):
        with ZipFile(stream, "a") as archive:
            archive.writestr(SLIDE, b"<bad")
    assert inspect(stream.getvalue()).exit_code == 3


def test_external_relationship_not_fetched() -> None:
    data = fixture()
    part = "ppt/_rels/presentation.xml.rels"
    node = root(data, part)
    for rel in node:
        if rel.get("Type", "").endswith("/slide"):
            rel.set("TargetMode", "External")
            rel.set("Target", "https://example.invalid/private.pptx")
    r = inspect(save(data, node, part))
    assert r.exit_code == 3
    assert any("external" in f["evidence"] for f in r.findings)


def test_ac18_source_qualified_identity() -> None:
    data = fixture(-1)
    slide = root(data)
    s = shape(slide)
    required(s, "p:nvSpPr/p:cNvPr", NS).set("id", "99")
    master = root(data, MASTER)
    required(master, "p:cSld/p:spTree", NS).append(deepcopy(s))
    data = save(save(data, slide), master, MASTER)
    r = inspect(data)
    fs = findings(r)
    assert len(fs) == 2 and len({f["fingerprint"] for f in fs}) == 2
    assert {f["shapes"][0]["origin"] for f in fs} == {"slide", "master"}


def test_ac19_text_inheritance() -> None:
    data = fixture()
    node = root(data)
    s = shape(node)
    p = required(s, "p:txBody/a:p", NS)
    ppr = required(p, "a:pPr", NS)
    ppr.attrib.update(marL="1200", indent="-100")
    spacing = ET.SubElement(ppr, A + "lnSpc")
    ET.SubElement(spacing, A + "spcPct", val="120000")
    run = ET.SubElement(p, A + "r")
    ET.SubElement(run, A + "t").text = "Hello"
    master = root(data, MASTER)
    default = required(master, "p:txStyles/p:otherStyle/a:lvl1pPr/a:defRPr", NS)
    default.set("sz", "2400")
    r = inspect(save(save(data, node), master, MASTER))
    props = r.shapes[0]["properties"]["paragraphs"][0]
    assert props["runs"][0]["properties"]["sz"] == "2400"
    assert props["runs"][0]["properties"]["latin"]
    assert props["properties"]["marL"] == "1200"
    assert props["children"]["lnSpc"]["children"][0]["attributes"]["val"] == "120000"
    assert not r.unknown


def test_table_cells_are_candidates() -> None:
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    slide.shapes.add_table(2, 2, Emu(-1), Emu(0), Emu(100), Emu(100))
    stream = BytesIO()
    p.save(stream)
    r = inspect(stream.getvalue())
    assert r.to_dict()["coverage"]["off-slide"]["candidates"] == 5
    cells = [s for s in r.shapes if s["shapes"][0]["kind"] == "table-cell"]
    assert len(cells) == 4
    assert len(findings(r)) == 3


def test_unknown_property_is_visible() -> None:
    data = fixture()
    node = root(data)
    ET.SubElement(required(shape(node), "p:spPr", NS), A + "futureProperty")
    r = inspect(save(data, node))
    assert r.exit_code == 2
    assert any("futureProperty" in u["reason"] for u in r.unknown)


def inspect(source, *, config=None, bounds_provider=None):
    """Keep foundation assertions scoped to the original three rules."""
    from dataclasses import replace

    config = config or Config()
    if isinstance(config, Config):
        config = replace(
            config, suppress=[*config.suppress, *({"rule": r} for r in NEW_RULES)]
        )
    return inspect_all(source, config=config, bounds_provider=bounds_provider)
