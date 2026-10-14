"""テスト実行中に PPTX と壊れた入力を作るための最小部品。"""

from io import BytesIO
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu


def single_shape_pptx(*, x: int, y: int, width: int, height: int) -> bytes:
    """位置と大きさを EMU 整数で指定した図形を 1 個作る。"""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(width), Emu(height))
    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def replace_part(package: bytes, name: str, content: bytes | None) -> bytes:
    """既存部品を差し替える。content が None なら削除する。"""
    output = BytesIO()
    with ZipFile(BytesIO(package)) as source, ZipFile(output, "w") as target:
        if name not in source.namelist():
            raise KeyError(name)
        for entry in source.infolist():
            if entry.filename == name:
                if content is not None:
                    target.writestr(entry, content)
            else:
                target.writestr(entry, source.read(entry))
    return output.getvalue()


def truncate_zip(package: bytes, offset: int) -> bytes:
    """ZIP の指定位置以降を切り落とす。"""
    if not 0 <= offset < len(package):
        raise ValueError("切断位置は ZIP のバイト列内で指定してください")
    return package[:offset]
