"""Issue #6 saved paint/image fixtures and independent candidate checks."""

import json
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from jsonschema import Draft202012Validator
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from pptviper import Config, inspect
from pptviper.backgrounds import behind, paint_colors
from pptviper.color_rules import nontext_contrast, text_contrast
from pptviper.geometry import A
from pptviper.rule_catalog import COLOR_RULES
from tests.pptx_factory import replace_part
from tests.test_static_rules import font, item


def color(rgb, direct=False, alpha: float = 1):
    return dict(
        rgb=rgb, alpha=alpha, source=dict(tag="srgbClr" if direct else "schemeClr")
    )


def fill(rgb, alpha: float = 1):
    return dict(kind="solidFill", colors=[color(rgb, alpha=alpha)], unknown=[])


def context() -> dict[str, Any]:
    return dict(
        size=[10000000, 7500000],
        background={**fill([255, 255, 255]), "source": "master"},
    )


def painted(
    key="a", text="", rgb=(0, 0, 0), order=0, box=(1000000, 1000000, 2000000, 2000000)
):
    i = item(key, box, text)
    i.props["visual"] = dict(fill=fill(rgb), line={}, order=order)
    font(i)["color"] = color([255, 255, 255])
    return i


def save(p):
    out = BytesIO()
    p.save(out)
    return out.getvalue()


def findings(r, rule):
    return [f for f in r.findings if f["rule"] == rule]


def run(data, **settings):
    return inspect(
        data, config=Config(measure_text=False, system_fonts=False, **settings)
    )


def text_deck(colors, theme=False):
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    sh = s.shapes.add_textbox(Emu(1000000), Emu(1000000), Emu(2000000), Emu(1000000))
    paragraph = sh.text_frame.paragraphs[0]
    for rgb in colors:
        r = paragraph.add_run()
        r.text = "Text"
        r.font.size = Pt(12)
        if theme:
            r.font.color.theme_color = MSO_THEME_COLOR.TEXT_1
        else:
            r.font.color.rgb = RGBColor(*rgb)
    return save(p)


def image_deck(kind="PNG", delta=0, resolution=300, crop: float = 0, text=False):
    image = Image.new("RGB", (resolution, resolution), (128, 128, 128))
    stream = BytesIO()
    image.save(stream, format=kind)
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    pic = s.shapes.add_picture(
        BytesIO(stream.getvalue()),
        Emu(1000000),
        Emu(1000000),
        Emu(914400 + delta),
        Emu(914400),
    )
    pic.crop_left = crop
    if text:
        sh = s.shapes.add_textbox(Emu(1100000), Emu(1100000), Emu(400000), Emu(400000))
        sh.text = "Text"
        r = sh.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(255, 255, 255)
    return save(p)


@pytest.mark.parametrize("rule", COLOR_RULES)
@pytest.mark.parametrize("positive", [False, True])
def test_each_rule_positive_negative(rule, positive):
    if rule == "low-contrast":
        data = text_deck([(128, 128, 128) if positive else (0, 0, 0)])
    elif rule == "off-theme-color":
        data = text_deck([(0, 0, 0)], theme=not positive)
    elif rule == "near-duplicate-color":
        data = text_deck([(0, 0, 0), (1, 1, 1) if positive else (255, 255, 255)])
    elif rule == "color-count":
        data = text_deck([(k * 10, 0, 0) for k in range(15 if positive else 13)])
    elif rule == "low-contrast-non-text":
        p = Presentation()
        s = p.slides.add_slide(p.slide_layouts[6])
        sh = s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Emu(1000000), Emu(1000000), Emu(1000000), Emu(1000000)
        )
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor(*((240, 240, 240) if positive else (0, 0, 0)))
        sh.line.fill.background()
        data = save(p)
    elif rule == "image-distorted":
        data = image_deck(delta=int(positive))
    elif rule == "image-low-resolution":
        data = image_deck(resolution=149 if positive else 150)
    else:
        data = image_deck(crop=-0.01 if positive else 0)
    r = run(data)
    assert len(findings(r, rule)) == int(positive)


@pytest.mark.parametrize("kind", ["PNG", "JPEG", "GIF", "BMP"])
def test_one_emu_distortion_all_formats(kind):
    assert len(findings(run(image_deck(kind, delta=1)), "image-distorted")) == 1


@pytest.mark.parametrize("kind", ["PNG", "JPEG", "EMF"])
def test_image_background_never_claims_contrast(kind):
    data = image_deck("PNG" if kind == "EMF" else kind, text=True)
    if kind == "EMF":
        with ZipFile(BytesIO(data)) as z:
            media = next(n for n in z.namelist() if n.startswith("ppt/media/"))
        data = replace_part(data, media, b"opaque EMF fixture")
    r = run(data)
    assert not findings(r, "low-contrast")
    assert any(
        u["rule"] == "low-contrast" and "Image" in u["reason"] for u in r.unknown
    )


def test_alpha_and_partial_occlusion():
    back = painted("back", rgb=(0, 0, 0), order=0)
    front = painted("front", text="White text", rgb=(255, 255, 255), order=1)
    front.props["visual"]["fill"] = fill([255, 255, 255], 0.5)
    result = text_contrast(front, [back, front], context(), "AA")
    assert result.value is not None
    assert result.hit and float(result.value) == pytest.approx(3.976653, abs=1e-6)
    # A partial opaque foreground must not erase the uncovered black background.
    part = painted(
        "part", rgb=(255, 255, 255), order=1, box=(1000000, 1000000, 1500000, 2000000)
    )
    front.props["visual"]["order"] = 2
    front.props["visual"]["fill"] = {}
    samples = behind(
        front, [back, part, front], context()["background"], context()["size"]
    )
    assert {tuple(c["rgb"]) for c in samples.colors} == {(0, 0, 0), (255, 255, 255)}


def test_large_text_aaa_and_outline_choice():
    i = painted(text="Text", rgb=(128, 128, 128))
    font(i)["sz"] = "1800"
    assert not text_contrast(i, [i], context(), "AA").hit
    assert text_contrast(i, [i], context(), "AAA").hit
    i = painted(rgb=(255, 255, 255))
    i.props["visual"]["line"] = fill([0, 0, 0])
    assert not nontext_contrast(i, [i], context()).hit


def test_gradient_clipping_and_pattern_candidates():
    from fractions import Fraction as F

    paint = dict(
        kind="gradFill",
        angle="0",
        colors=[color([0, 0, 0]), color([255, 255, 255])],
        stops=[
            dict(position="0", color=color([0, 0, 0])),
            dict(position="100000", color=color([255, 255, 255])),
        ],
    )
    result = paint_colors(
        paint, (F(25), F(0), F(75), F(100)), (F(0), F(0), F(100), F(100))
    )
    assert [c["rgb"] for c in result.colors] == [[63.75] * 3, [191.25] * 3]
    paint = dict(kind="pattFill", colors=[color([1, 2, 3]), color([4, 5, 6])])
    assert (
        len(
            paint_colors(
                paint, (F(0), F(0), F(0), F(0)), (F(0), F(0), F(0), F(0))
            ).colors
        )
        == 2
    )


def test_master_background_and_schema():
    data = text_deck([(255, 255, 255)])
    from xml.etree import ElementTree as ET

    with ZipFile(BytesIO(data)) as z:
        root = ET.fromstring(z.read("ppt/slideMasters/slideMaster1.xml"))
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    bg = root.find("p:cSld/p:bg", ns)
    assert bg is not None
    bg.clear()
    pr = ET.SubElement(bg, "{" + ns["p"] + "}bgPr")
    solid = ET.SubElement(pr, A + "solidFill")
    ET.SubElement(solid, A + "srgbClr", val="000000")
    r = run(replace_part(data, "ppt/slideMasters/slideMaster1.xml", ET.tostring(root)))
    assert not findings(r, "low-contrast")
    assert r.slide_contexts["256"]["background"]["source"] == "master"
    schema = json.loads(
        (Path(__file__).parents[1] / "src/pptviper/report.schema.json").read_text()
    )
    Draft202012Validator(schema).validate(r.to_dict())
    for rule in COLOR_RULES:
        c = r.to_dict()["coverage"][rule]
        assert c["candidates"] == sum(
            c[s] for s in ("pass", "finding", "unknown", "not-applicable")
        )
    assert (
        r.to_json()
        == run(
            replace_part(data, "ppt/slideMasters/slideMaster1.xml", ET.tostring(root))
        ).to_json()
    )


def test_invalid_contrast_level():
    with pytest.raises(ValueError):
        Config(contrast_level="bad")


@pytest.mark.parametrize(
    "limit,level,size", [(4.5, "AA", "1200"), (3, "AA", "1800"), (7, "AAA", "1200")]
)
def test_threshold_and_one_channel_step(limit, level, size):
    linear = 1.05 / limit - 0.05
    channel = (1.055 * linear ** (1 / 2.4) - 0.055) * 255
    i = painted(text="Text", rgb=(channel,) * 3)
    font(i)["sz"] = size
    assert not text_contrast(i, [i], context(), level).hit
    i.props["visual"]["fill"] = fill([channel + 1] * 3)
    assert text_contrast(i, [i], context(), level).hit


def test_saved_alpha_composition():
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    for index, rgb in enumerate(((0, 0, 0), (255, 255, 255))):
        sh = s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Emu(1000000), Emu(1000000), Emu(2000000), Emu(1000000)
        )
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor(*rgb)
        sh.line.fill.background()
        if index:
            clr = sh._element.find(".//" + A + "solidFill/" + A + "srgbClr")
            assert clr is not None
            clr.append(clr.makeelement(A + "alpha", {"val": "50000"}))
            sh.text = "White"
            r = sh.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(255, 255, 255)
    r = run(save(p))
    result = findings(r, "low-contrast")
    assert len(result) == 1 and result[0]["measured"]["value"] == pytest.approx(
        3.976653, abs=1e-6
    )


def test_direct_color_representations():
    from xml.etree.ElementTree import Element

    from pptviper.inheritance import color as resolve_color

    for tag, attrs, expected in [
        ("prstClr", {"val": "dkBlue"}, [0, 0, 139]),
        ("scrgbClr", {"r": "0", "g": "100000", "b": "0"}, [0, 255, 0]),
        ("hslClr", {"hue": "0", "sat": "100000", "lum": "50000"}, [255, 0, 0]),
    ]:
        unknown = []
        result = resolve_color(Element(A + tag, attrs), {}, {}, unknown)
        assert result and result["rgb"] == expected and not unknown
