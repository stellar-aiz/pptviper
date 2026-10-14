"""実機証跡を捏造せず、検証ツールの拒否条件を人工データで試す。"""

import csv
import json
from copy import deepcopy
from fractions import Fraction
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from pptviper import Config, inspect
from pptviper.rule_catalog import RULES
from scripts.acceptance.compare import compare
from scripts.acceptance.generate import generate, write_json
from scripts.acceptance.review import aggregate, build, digest, seed_check


@pytest.fixture(scope="module")
def generated(tmp_path_factory):
    folder = tmp_path_factory.mktemp("acceptance")
    manifest = generate(folder)
    return folder, manifest


def test_deterministic_generation(generated, tmp_path):
    folder, manifest = generated
    assert len(manifest["cases"]) == 216
    generate(tmp_path)
    for name in ("calibration.pptx", "calibration.json", "seeded.pptx", "seeded.json"):
        assert (folder / name).read_bytes() == (tmp_path / name).read_bytes()
    assert len(Presentation(str(folder / "seeded.pptx")).slides) == 10


def test_seed_truth_and_image_unknown(generated):
    folder, _ = generated
    report = inspect(
        folder / "seeded.pptx", config=Config(system_fonts=False)
    ).to_dict()
    truth = json.loads((folder / "seeded.json").read_text())
    assert seed_check(report, truth, digest(folder / "seeded.pptx"))["passed"]
    report["unknown"] = [u for u in report["unknown"] if u["rule"] != "low-contrast"]
    assert not seed_check(report, truth, digest(folder / "seeded.pptx"))["passed"]


def write_csv(path, rows):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=["name", "slide", "shape_id", "lines", "width_pt", "height_pt"],
        )
        writer.writeheader()
        writer.writerows(rows)


@pytest.fixture(scope="module")
def measurements(generated):
    folder, manifest = generated
    report = inspect(
        folder / "calibration.pptx", config=Config(system_fonts=False)
    ).to_dict()
    by_name = {
        s["shapes"][0]["name"]: s["properties"].get("text_layout", {})
        for s in report["shapes"]
        if s["shapes"][0]["origin"] == "slide"
    }
    return [
        dict(
            name=c["name"],
            slide=c["slide"],
            shape_id=c["shape_id"],
            lines=by_name[c["name"]]["lines"],
            width_pt=str(Fraction(by_name[c["name"]]["width"]) / 12700),
            height_pt=str(Fraction(by_name[c["name"]]["height"]) / 12700),
        )
        for c in manifest["cases"]
    ]


def test_underestimate_is_failure(generated, measurements, tmp_path):
    folder, manifest = generated
    path = tmp_path / "synthetic.csv"
    rows = deepcopy(measurements)
    rows[0]["lines"] += 1
    write_csv(path, rows)
    result = compare(
        folder / "calibration.pptx", manifest, path, config=Config(system_fonts=False)
    )
    assert not result["passed"]
    assert [r["name"] for r in result["results"] if not r["passed"]] == [
        rows[0]["name"]
    ]
    assert result["results"][0]["delta"][0] == "-1"


@pytest.mark.parametrize(
    "kind", ["missing", "duplicate", "wrong-id", "nan", "too-small"]
)
def test_invalid_measurements_rejected(generated, measurements, tmp_path, kind):
    folder, manifest = generated
    rows = deepcopy(measurements)
    mode = "macro"
    if kind == "missing":
        rows.pop()
    elif kind == "duplicate":
        rows[-1] = rows[0]
    elif kind == "wrong-id":
        rows[0]["shape_id"] = 999
    elif kind == "nan":
        rows[0]["width_pt"] = "NaN"
    else:
        rows = rows[:49]
        mode = "manual"
    path = tmp_path / "invalid.csv"
    write_csv(path, rows)
    with pytest.raises(ValueError):
        compare(
            folder / "calibration.pptx",
            manifest,
            path,
            mode,
            Config(system_fonts=False),
        )


def test_manual_subset_is_explicit(generated, measurements, tmp_path):
    folder, manifest = generated
    path = tmp_path / "manual.csv"
    write_csv(path, measurements[:50])
    result = compare(
        folder / "calibration.pptx",
        manifest,
        path,
        "manual",
        Config(system_fonts=False),
    )
    assert result["passed"] and result["measured"] == 50
    assert len(result["omitted"]) == 166


def test_review_html_binds_evidence(generated, tmp_path):
    folder, _ = generated
    for n in range(1, 11):
        Image.new("RGB", (400, 300), (255, 255, 255)).save(tmp_path / f"{n}.png")
    output = tmp_path / "review.html"
    payload = build(folder / "seeded.pptx", tmp_path, output, "generated")
    assert payload["deck_sha256"] == digest(folder / "seeded.pptx")
    assert len(payload["slides"]) == 10 and len(payload["rules"]) == 44
    html = output.read_text()
    assert "__PAYLOAD__" not in html and "data:image/png;base64," in html
    assert "checked_rules" in html and "report_sha256" in html


def test_aggregate_requires_real_corpus():
    with pytest.raises(ValueError, match="3 本"):
        aggregate([], Path("."))


def test_aggregate_accepts_complete_synthetic_contract_and_rejects_missing_review(
    generated, tmp_path
):
    # 人の証跡ではなく、集計契約だけを検証する fixture。
    folder, _ = generated
    entries = []
    for i in range(3):
        p = Presentation(str(folder / "seeded.pptx"))
        p.core_properties.title = f"synthetic-{i}"
        deck = tmp_path / f"deck-{i}.pptx"
        p.save(deck)
        report = inspect(deck, config=Config(system_fonts=False)).to_dict()
        report_path = tmp_path / f"report-{i}.json"
        write_json(report_path, report)
        images = tmp_path / f"images-{i}"
        images.mkdir()
        image_refs = []
        for n in range(1, 11):
            Image.new("RGB", (40, 30)).save(images / f"{n}.png")
            image_refs.append(dict(number=n, sha256=digest(images / f"{n}.png")))
        origin = "generated" if i == 0 else "human"
        review = dict(
            deck_sha256=digest(deck),
            report_sha256=digest(report_path),
            reviewed_slides=list(range(1, 11)),
            checked_rules=list(RULES),
            operator="synthetic test",
            os="synthetic",
            powerpoint_version="synthetic",
            confidentiality_reviewed=True,
            provenance=origin,
            images=image_refs,
            false_positives=[],
            misses=[],
        )
        write_json(tmp_path / f"review-{i}.json", review)
        entry = dict(
            deck=deck.name,
            report=report_path.name,
            review=f"review-{i}.json",
            images=images.name,
            provenance=origin,
        )
        if i == 0:
            truth = json.loads((folder / "seeded.json").read_text())
            truth["deck_sha256"] = digest(deck)
            write_json(tmp_path / "truth.json", truth)
            entry["truth"] = "truth.json"
        entries.append(entry)
    assert aggregate(entries, tmp_path)["passed"]
    review = json.loads((tmp_path / "review-2.json").read_text())
    review["reviewed_slides"].pop()
    write_json(tmp_path / "review-2.json", review)
    with pytest.raises(ValueError, match="レビュー漏れ"):
        aggregate(entries, tmp_path)
