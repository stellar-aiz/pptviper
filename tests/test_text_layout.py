"""Text layout acceptance fixtures use a generated MIT-licensed test font."""

from fractions import Fraction as F
from io import BytesIO
from pathlib import Path
from typing import Any, cast

import pytest
from fontTools.fontBuilder import FontBuilder
from fontTools.pens.ttGlyphPen import TTGlyphPen
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Emu, Pt

from pptviper import Config, inspect
from pptviper.fonts import FontBook
from pptviper.linebreak import Glyph, wrap
from pptviper.text_layout import layout


@pytest.fixture
def font_path(tmp_path: Path) -> Path:
    """Original synthetic glyphs under the repository's MIT license, never shipped."""
    builder = FontBuilder(1000, isTTF=True)
    builder.setupGlyphOrder([".notdef", "box"])
    builder.setupCharacterMap(
        {
            ord(ch): "box"
            for ch in (
                "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz"
                "0123456789 -.,()漢字。、「」ぁ"
            )
        }
    )
    glyphs = {}
    for name in [".notdef", "box"]:
        pen = TTGlyphPen(None)
        pen.moveTo((0, 0))
        pen.lineTo((500, 0))
        pen.lineTo((500, 700))
        pen.lineTo((0, 700))
        pen.closePath()
        glyphs[name] = pen.glyph()
    builder.setupGlyf(glyphs)
    builder.setupHorizontalMetrics({name: (1000, 0) for name in glyphs})
    builder.setupHorizontalHeader(ascent=800, descent=-200)
    builder.setupNameTable(
        dict(
            familyName="Pptviper Fixture",
            styleName="Regular",
            uniqueFontIdentifier="Pptviper Fixture Regular",
            fullName="Pptviper Fixture Regular",
            psName="PptviperFixture-Regular",
            version="Version 1.0",
            licenseDescription="Original test font; MIT; see repository LICENSE.",
        )
    )
    builder.font["name"].setName("検証書体", 1, 3, 1, 0x411)
    builder.setupOS2(
        sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200
    )
    builder.setupPost()
    builder.setupMaxp()
    path = tmp_path / "fixture.ttf"
    builder.save(path)
    return path


def props(text="AA", **body) -> dict[str, Any]:
    return dict(
        body=dict(lIns="0", rIns="0", tIns="0", bIns="0", **body),
        autofit=dict(mode="noAutofit"),
        paragraphs=[
            dict(
                properties={},
                children={},
                default_run={
                    "sz": "1200",
                    "latin": "Pptviper Fixture",
                    "ea": "Pptviper Fixture",
                },
                runs=[
                    dict(
                        text=text,
                        kind="r",
                        properties={
                            "sz": "1200",
                            "latin": "Pptviper Fixture",
                            "ea": "Pptviper Fixture",
                        },
                    )
                ],
            )
        ],
    )


def pptx(
    text="AAA",
    width=304800,
    height=152400,
    fit=MSO_AUTO_SIZE.NONE,
    word_wrap=False,
    second=False,
) -> bytes:
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    shape = s.shapes.add_textbox(Emu(1000000), Emu(1000000), Emu(width), Emu(height))
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    tf.word_wrap = word_wrap
    tf.auto_size = fit
    shape.text = text
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(12)
    run.font.name = "Pptviper Fixture"
    if second:
        s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Emu(1000000), Emu(1200000), Emu(304800), Emu(100000)
        )
    stream = BytesIO()
    p.save(stream)
    return stream.getvalue()


def matches(report, rule):
    return [f for f in report.findings if f["rule"] == rule]


def config(path):
    return Config(font_paths=[str(path)], system_fonts=False)


def test_exact_width_and_one_character_overflow(font_path):
    book = FontBook([str(font_path)], False)
    assert (
        layout(props("AA", wrap="none"), F(304800), F(152400), book)["overflow_width"]
        == "0"
    )
    assert (
        layout(props("AAA", wrap="none"), F(304800), F(152400), book)["overflow_width"]
        == "152400"
    )
    for text, count in [("AA", 0), ("AAA", 1)]:
        report = inspect(pptx(text=text), config=config(font_path))
        assert len(matches(report, "text-overflow")) == count
        assert not matches(report, "font-missing")


def test_localized_names_and_missing_glyph(font_path):
    book = FontBook([str(font_path)], False)
    face, data = book.select("検証書体", False, False)
    assert face and data and data.units == 1000
    width, _, unknown, missing = book.measure("☃", {"latin": "Pptviper Fixture"}, F(12))
    assert width == 152400 and unknown and missing is None
    width, _, unknown, _ = book.measure(
        "A", {"latin": "Pptviper Fixture", "b": "1"}, F(12)
    )
    assert width == F(152400) * F(17, 16) and any(
        "Synthetic bold" in r for r in unknown
    )


def test_missing_font_still_measures():
    r = inspect(pptx(), config=Config(system_fonts=False))
    assert (
        len(matches(r, "font-missing")) == 1 and len(matches(r, "text-overflow")) == 1
    )
    assert matches(r, "text-overflow")[0]["confidence"] == "worst-case"
    assert "1 em" in matches(r, "text-overflow")[0]["evidence"]


def test_japanese_prohibitions():
    glyphs = [Glyph(ch, F(1), F(1)) for ch in "漢字。漢字"]
    lines, _ = wrap(glyphs, F(2), F(4), True)
    strings = ["".join(g.text for g in line) for line in lines]
    assert strings == ["漢", "字。", "漢字"]
    lines, _ = wrap([Glyph(ch, F(1), F(1)) for ch in "漢「字"], F(2), F(4), True)
    assert ["".join(g.text for g in line) for line in lines] == ["漢", "「字"]


def test_wrapping_rules_positive_negative(font_path):
    for text, expect in [("ABCDE", 1), ("AB", 0)]:
        r = inspect(
            pptx(text=text, word_wrap=True, height=1000000), config=config(font_path)
        )
        assert len(matches(r, "orphan-line")) == expect
    r = inspect(
        pptx(text="ABCDE", word_wrap=True, height=1000000), config=config(font_path)
    )
    assert (
        len(matches(r, "unbreakable-word")) == 1
        and len(matches(r, "excessive-wrap")) == 1
    )
    r = inspect(
        pptx(text="A", word_wrap=True, height=1000000), config=config(font_path)
    )
    assert not matches(r, "unbreakable-word") and not matches(r, "excessive-wrap")


def test_autofit_expansion_reaches_overlap_and_offslide(font_path):
    data = pptx(
        text="ABCDEFGH",
        word_wrap=True,
        fit=MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT,
        second=True,
    )
    r = inspect(data, config=config(font_path))
    assert len(matches(r, "overlap")) == 1
    assert matches(r, "overlap")[0]["confidence"] == "worst-case"
    measured = r.shapes[0]["properties"]["text_layout"]
    assert measured["lines"] == 4 and measured["expand_height"]
    p = Presentation(BytesIO(data))
    p.slides[0].shapes[0].top = Emu(int(p.slide_height or 0) - 152400)
    stream = BytesIO()
    p.save(stream)
    assert matches(inspect(stream.getvalue(), config=config(font_path)), "off-slide")


def test_vertical_text_unknown(font_path):
    p = Presentation(BytesIO(pptx()))
    cast(Any, p.slides[0].shapes[0]).text_frame._txBody.bodyPr.set("vert", "eaVert")
    stream = BytesIO()
    p.save(stream)
    r = inspect(stream.getvalue(), config=config(font_path))
    assert not matches(r, "text-overflow")
    assert any(
        u["rule"] == "text-overflow" and "Vertical" in u["reason"] for u in r.unknown
    )


def test_norm_autofit_missing_scale_and_tabs(font_path):
    book = FontBook([str(font_path)], False)
    p = props("A\tB", wrap="none")
    p["autofit"] = {"mode": "normAutofit", "fontScale_defaulted": True}
    result = layout(p, F(1000000), F(1000000), book)
    assert result["width"] == "609600" and any(
        "no shrink" in r for r in result["unknown"]
    )
    p = props("AA")
    p["autofit"] = {"mode": "normAutofit", "fontScale": "50000"}
    assert layout(p, F(1000000), F(1000000), book)["width"] == "152400"


def test_tables_growth_and_cell_overflow(font_path):
    def deck(text):
        p = Presentation()
        s = p.slides.add_slide(p.slide_layouts[6])
        shape = s.shapes.add_table(
            1,
            1,
            Emu(1000000),
            Emu(int(p.slide_height or 0) - 300000),
            Emu(400000),
            Emu(300000),
        )
        cell = shape.table.cell(0, 0)
        cell.text = text
        cell.text_frame.word_wrap = True
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.name = "Pptviper Fixture"
        run.font.size = Pt(12)
        stream = BytesIO()
        p.save(stream)
        return stream.getvalue()

    r = inspect(deck("ABCDEFGHIJKLMN"), config=config(font_path))
    assert (
        len(matches(r, "table-cell-overflow")) == 1
        and len(matches(r, "table-overflow")) == 1
    )
    r = inspect(deck("A"), config=config(font_path))
    assert not matches(r, "table-cell-overflow") and not matches(r, "table-overflow")


def test_character_properties_and_spacing(font_path):
    p = props("aa", wrap="none")
    p["paragraphs"][0]["runs"][0]["properties"].update(
        cap="small", spc="100", baseline="30000"
    )
    result = layout(p, F(1000000), F(1000000), FontBook([str(font_path)], False))
    assert F(result["width"]) == 2 * (F(152400) * F(4, 5) + 12700)
    assert any("Baseline" in r for r in result["unknown"])


def test_layout_coverage_and_schema(font_path):
    import json

    from jsonschema import Draft202012Validator

    from pptviper.rule_catalog import LAYOUT_RULES

    r = inspect(pptx(), config=config(font_path))
    data = r.to_dict()
    for rule in LAYOUT_RULES:
        c = data["coverage"][rule]
        assert c["candidates"] == sum(
            c[k] for k in ["pass", "finding", "unknown", "not-applicable"]
        )
    schema = json.loads(
        (Path(__file__).parents[1] / "src/pptviper/report.schema.json").read_text()
    )
    Draft202012Validator(schema).validate(data)
    assert r.to_json() == inspect(pptx(), config=config(font_path)).to_json()


def test_bullets_consume_width_and_known_font_coverage(font_path):
    p = props("AAA")
    p["paragraphs"][0]["children"]["buChar"] = {"attributes": {"char": "A"}}
    result = layout(p, F(457200), F(1000000), FontBook([str(font_path)], False))
    assert result["lines"] == 2 and result["fonts"] == ["Pptviper Fixture"]
    r = inspect(pptx(text="A"), config=config(font_path))
    assert r.to_dict()["coverage"]["font-missing"]["pass"] == 1
