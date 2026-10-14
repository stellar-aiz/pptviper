"""Check the minimal CLI, PPTX fixtures, and network blocking."""

import socket
from importlib.metadata import version
from io import BytesIO
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

import pytest
from pptx import Presentation
from pytest_socket import SocketBlockedError

from pptviper import __version__
from pptviper.cli import main
from tests.pptx_factory import replace_part, single_shape_pptx, truncate_zip


def test_version(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as result:
        main(["--version"])
    assert result.value.code == 0
    assert capsys.readouterr().out.strip() == f"pptviper {version('pptviper')}"
    assert __version__ == version("pptviper")


def test_cli_does_not_pretend_to_inspect() -> None:
    with pytest.raises(SystemExit) as result:
        main(["example.pptx"])
    assert result.value.code == 2


@pytest.fixture
def package() -> bytes:
    return single_shape_pptx(x=-1, y=360001, width=720001, height=360001)


def test_shape_coordinates(package: bytes) -> None:
    with ZipFile(BytesIO(package)) as archive:
        slide = ElementTree.fromstring(archive.read("ppt/slides/slide1.xml"))
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    offset = slide.find(".//a:xfrm/a:off", ns)
    extent = slide.find(".//a:xfrm/a:ext", ns)
    assert offset is not None and extent is not None
    assert offset.attrib == {"x": "-1", "y": "360001"}
    assert extent.attrib == {"cx": "720001", "cy": "360001"}
    presentation = Presentation(BytesIO(package))
    assert len(presentation.slides) == 1
    assert len(presentation.slides[0].shapes) == 1


def test_replace_and_delete_part(package: bytes) -> None:
    name = "ppt/slides/slide1.xml"
    changed = replace_part(package, name, b"<broken")
    with ZipFile(BytesIO(package)) as original, ZipFile(BytesIO(changed)) as archive:
        assert archive.read(name) == b"<broken"
        for part in original.namelist():
            if part != name:
                assert archive.read(part) == original.read(part)
    with ZipFile(BytesIO(replace_part(package, name, None))) as archive:
        assert name not in archive.namelist()
    with pytest.raises(KeyError):
        replace_part(package, "missing.xml", b"x")


def test_truncated_zip(package: bytes) -> None:
    cut = truncate_zip(package, 10)
    assert cut == package[:10]
    with pytest.raises(BadZipFile):
        ZipFile(BytesIO(cut))
    with pytest.raises(ValueError):
        truncate_zip(package, -1)


def test_network_is_blocked() -> None:
    with pytest.raises(SocketBlockedError):
        socket.create_connection(("127.0.0.1", 9), timeout=0.1)
