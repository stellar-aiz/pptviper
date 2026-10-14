"""スライド画像・検査結果・人の注記を関連付ける。"""

import base64
import hashlib
import json
from collections import Counter
from pathlib import Path

from pptx import Presentation

from pptviper import inspect
from pptviper.rule_catalog import RULES
from scripts.acceptance.generate import write_json


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build(deck: Path, images: Path, output: Path, provenance: str) -> dict:
    if provenance not in ("generated", "human"):
        raise ValueError("作成元を generated または human で指定してください")
    presentation = Presentation(str(deck))
    if presentation.slide_width is None or presentation.slide_height is None:
        raise ValueError("スライドのサイズがありません")
    report = inspect(deck).to_dict()
    output.parent.mkdir(parents=True, exist_ok=True)
    report_path = output.with_suffix(".report.json")
    write_json(report_path, report)
    slides = []
    for number in range(1, len(presentation.slides) + 1):
        path = images / f"{number}.png"
        raw = path.read_bytes()
        if not raw.startswith(b"\x89PNG\r\n\x1a\n"):
            raise ValueError(f"PNG ではありません: {path}")
        slides.append(
            dict(
                number=number,
                image="data:image/png;base64," + base64.b64encode(raw).decode(),
                sha256=digest(path),
            )
        )
    payload = dict(
        schema_version=1,
        deck_sha256=digest(deck),
        report_sha256=digest(report_path),
        provenance=provenance,
        width=int(presentation.slide_width),
        height=int(presentation.slide_height),
        slides=slides,
        report=report,
        rules=list(RULES),
    )
    serialized = json.dumps(payload, ensure_ascii=True).replace("<", "\\u003c")
    template = Path(__file__).with_name("review.html").read_text(encoding="utf-8")
    output.write_text(template.replace("__PAYLOAD__", serialized), encoding="utf-8")
    return payload


def seed_check(report: dict, truth: dict, deck_hash: str) -> dict:
    if truth["deck_sha256"] != deck_hash:
        raise ValueError("正解データとデッキが一致しません")
    missing = []
    for expected in truth["expected"]:
        records = (
            report["unknown"] if expected["kind"] == "unknown" else report["findings"]
        )
        found = any(
            r["rule"] == expected["rule"]
            and (r.get("slide") or {}).get("number") == expected["slide"]
            and any(
                str(s["id"]) == str(expected["shape_id"]) and s.get("origin") == "slide"
                for s in r["shapes"]
            )
            and not r.get("suppressed", False)
            for r in records
        )
        if not found:
            missing.append(expected)
    return dict(
        passed=bool(truth["expected"]) and not missing,
        missing=missing,
        checked=len(truth["expected"]),
    )


def aggregate(entries: list[dict], root: Path) -> dict:
    if len(entries) < 3:
        raise ValueError("10 枚以上のデッキが 3 本以上必要です")
    false_positives = Counter()
    missed = []
    decks = set()
    human = 0
    seeded = []
    for entry in entries:
        deck = root / entry["deck"]
        report_file = root / entry["report"]
        report = json.loads(report_file.read_text(encoding="utf-8"))
        review = json.loads((root / entry["review"]).read_text(encoding="utf-8"))
        sha = digest(deck)
        if (
            sha in decks
            or review["deck_sha256"] != sha
            or review["report_sha256"] != digest(report_file)
        ):
            raise ValueError("重複デッキまたは証跡のハッシュ不一致です")
        decks.add(sha)
        count = len(Presentation(str(deck)).slides)
        if count < 10 or sorted(review["reviewed_slides"]) != list(range(1, count + 1)):
            raise ValueError("スライド数不足またはレビュー漏れです")
        if set(review["checked_rules"]) != set(RULES):
            raise ValueError("確認観点のチェックが不足しています")
        for field in ("operator", "os", "powerpoint_version"):
            if not isinstance(review.get(field), str) or not review[field].strip():
                raise ValueError("実施者・OS・PowerPoint の版を記録してください")
        if review.get("confidentiality_reviewed") is not True:
            raise ValueError("機密確認の証跡がありません")
        if review.get("provenance") != entry["provenance"] or entry[
            "provenance"
        ] not in ("human", "generated"):
            raise ValueError("デッキの作成元が一致しません")
        human += entry["provenance"] == "human"
        images = review["images"]
        if len(images) != count or [i["number"] for i in images] != list(
            range(1, count + 1)
        ):
            raise ValueError("画像証跡が不足しています")
        for image in images:
            path = root / entry["images"] / f"{image['number']}.png"
            if digest(path) != image["sha256"]:
                raise ValueError("画像のハッシュが一致しません")
        by_id = {f["fingerprint"]: f for f in report["findings"]}
        ids = review["false_positives"]
        if len(ids) != len(set(ids)) or not set(ids) <= by_id.keys():
            raise ValueError("誤検出の指摘 ID が不正です")
        false_positives.update(by_id[i]["rule"] for i in ids)
        for miss in review["misses"]:
            if miss["slide"] not in range(1, count + 1) or not all(
                0 <= miss[k] <= 1 for k in ("x", "y")
            ):
                raise ValueError("見逃しの位置が不正です")
            if miss["rule"] not in RULES and not miss.get("note", "").strip():
                raise ValueError("規則名または自由記述が必要です")
            missed.append(dict(deck=entry["deck"], **miss))
        if "truth" in entry:
            truth = json.loads((root / entry["truth"]).read_text(encoding="utf-8"))
            if entry["provenance"] != "generated":
                raise ValueError("正解データは生成デッキに対応させてください")
            seeded.append(seed_check(report, truth, sha))
        if report.get("errors") or report.get("unenumerated"):
            missed.append(dict(deck=entry["deck"], note="解析が不完全です"))
    if human < 2 or not seeded:
        raise ValueError("人が作成したデッキ 2 本と、生成した正解付きデッキが必要です")
    return dict(
        passed=not missed and all(s["passed"] for s in seeded),
        missed=missed,
        seeded=seeded,
        false_positives=dict(sorted(false_positives.items())),
        decks=len(decks),
        human_decks=human,
    )
