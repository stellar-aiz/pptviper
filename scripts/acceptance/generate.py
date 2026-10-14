"""再現可能な較正デッキと既知問題デッキを生成する。"""

import hashlib
import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_STORED, ZipFile, ZipInfo

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def write_json(path: Path, value) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def save(presentation, path: Path) -> str:
    fixed = datetime(2000, 1, 1, tzinfo=UTC)
    presentation.core_properties.created = fixed
    presentation.core_properties.modified = fixed
    presentation.core_properties.author = "pptviper calibration generator"
    stream = BytesIO()
    presentation.save(stream)
    with ZipFile(stream) as source, ZipFile(path, "w") as target:
        for name in sorted(source.namelist()):
            entry = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_STORED
            entry.create_system = 0
            target.writestr(entry, source.read(name))
    return hashlib.sha256(path.read_bytes()).hexdigest()


def generate(directory: Path) -> dict:
    directory.mkdir(parents=True, exist_ok=True)
    p = Presentation()
    p.slide_width, p.slide_height = Inches(10), Inches(7.5)
    samples = (
        "Lightweight inspection: punctuation, words, and layout.",
        "生成した文章です。「括弧」と、句読点。折り返しを確認します。",
        "日本語 Mixed text 12345（layout）と English words.",
    )
    fonts = ("Arial", "Calibri", "Yu Gothic", "Meiryo", "Helvetica", "Hiragino Sans")
    cases = []
    for family in fonts:
        for language, text in enumerate(samples):
            for variant in range(12):
                index = len(cases)
                if index % 6 == 0:
                    slide = p.slides.add_slide(p.slide_layouts[6])
                slide = p.slides[-1]
                shape = slide.shapes.add_textbox(
                    Inches(0.4 + (index % 2) * 4.8),
                    Inches(0.3 + ((index % 6) // 2) * 2.3),
                    Inches(3.3 if variant % 2 else 2.1),
                    Inches(1.8),
                )
                shape.name = f"CAL-{index + 1:04d}"
                frame = shape.text_frame
                frame.word_wrap = variant != 10
                frame.auto_size = (
                    MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                    if variant == 8
                    else MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
                    if variant == 9
                    else MSO_AUTO_SIZE.NONE
                )
                para = frame.paragraphs[0]
                run = para.add_run()
                run.text = text + (
                    " Supercalifragilisticexpialidocious" * 3 if variant == 7 else ""
                )
                run.font.name = family
                run.font.size = Pt(12 + variant * 2)
                run.font.bold = variant in (1, 3)
                run.font.italic = variant in (2, 3)
                props = run._r.get_or_add_rPr()
                for tag in ("ea", "cs"):
                    element = props.makeelement(A + tag, {"typeface": family})
                    props.append(element)
                if variant in (4, 5):
                    props.set("cap", "all" if variant == 4 else "small")
                if variant == 6:
                    props.set("spc", "150")
                if variant == 11:
                    pr = para._p.get_or_add_pPr()
                    pr.set("marL", "254000")
                    pr.set("indent", "-127000")
                    pr.append(pr.makeelement(A + "buChar", {"char": "•"}))
                cases.append(
                    dict(
                        name=shape.name,
                        slide=len(p.slides),
                        shape_id=shape.shape_id,
                        font=family,
                        language=language,
                        variant=variant,
                    )
                )
    digest = save(p, directory / "calibration.pptx")
    manifest = dict(
        schema_version=1,
        deck_sha256=digest,
        cases=cases,
        provenance="generated",
        generator="scripts.acceptance.generate",
    )
    write_json(directory / "calibration.json", manifest)
    seeded = Presentation()
    expected = []
    for index in range(10):
        s = seeded.slides.add_slide(seeded.slide_layouts[6])
        shape = s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(-0.4), Inches(1), Inches(2), Inches(1)
        )
        shape.name = f"SEED-{index + 1:02d}"
        expected.append(
            dict(
                slide=index + 1,
                shape_id=shape.shape_id,
                rule="off-slide",
                kind="finding",
            )
        )
        if index % 2 == 0:
            image = BytesIO()
            Image.new("RGB", (100, 100), (40, 80, 120)).save(image, format="PNG")
            image.seek(0)
            s.shapes.add_picture(image, Inches(3), Inches(1), Inches(3), Inches(2))
            text_shape = s.shapes.add_textbox(
                Inches(3.2), Inches(1.2), Inches(2), Inches(1)
            )
            text_shape.text = "Image background"
            expected.append(
                dict(
                    slide=index + 1,
                    shape_id=text_shape.shape_id,
                    rule="low-contrast",
                    kind="unknown",
                )
            )
    seed_hash = save(seeded, directory / "seeded.pptx")
    write_json(
        directory / "seeded.json",
        dict(
            schema_version=1,
            deck_sha256=seed_hash,
            slides=10,
            provenance="generated",
            expected=expected,
        ),
    )
    return manifest
